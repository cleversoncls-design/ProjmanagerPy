"""Consumo já apropriado no sistema anterior (pedido do usuário).

Projetos migrados chegam com horas já consumidas a um custo médio por hora.
Cada lançamento (ProjectLegacyConsumption) soma ao custo real/margem do
projeto e às horas consumidas — sem virar apontamento (Timesheet).

- GET/POST   /projects/{id}/legacy-consumption
- PATCH/DELETE /projects/{id}/legacy-consumption/{entry_id}
- POST /legacy-consumption/import  (planilha .xlsx, vários projetos de uma vez)
- GET  /legacy-consumption/template.xlsx  (modelo da planilha)

Tudo restrito a MANAGEMENT_ROLES (dado financeiro)."""
from __future__ import annotations

import unicodedata
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from fastapi import APIRouter, Depends, File, HTTPException, Response, UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.datetime import from_excel
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..audit import record_audit
from ..database import get_db
from ..deps import MANAGEMENT_ROLES, require_project_access, require_roles
from ..i18n import t as translate
from ..models import AuditAction, Project, ProjectLegacyConsumption, User
from ..schemas import (
    LegacyConsumptionImportResult,
    ProjectLegacyConsumptionCreate,
    ProjectLegacyConsumptionRead,
    ProjectLegacyConsumptionUpdate,
)

router = APIRouter(tags=["legacy-consumption"])

MAX_IMPORT_BYTES = 2 * 1024 * 1024
MAX_IMPORT_ROWS = 5000
MAX_REPORTED_ERRORS = 15
_CENT = Decimal("0.01")


def _read(entry: ProjectLegacyConsumption) -> dict:
    return {
        "id": entry.id,
        "project_id": entry.project_id,
        "reference_date": entry.reference_date,
        "hours": entry.hours,
        "cost_per_hour": entry.cost_per_hour,
        "total_cost": (Decimal(entry.hours) * Decimal(entry.cost_per_hour)).quantize(_CENT),
        "description": entry.description,
    }


def _project_or_404(db: Session, project_id: str, user: User) -> Project:
    project = db.get(Project, project_id)
    if not project:
        raise HTTPException(status_code=404, detail=translate("Projeto não encontrado", user.language))
    require_project_access(project, user, write=True, allow_consultant_write=False)
    return project


def _entry_or_404(db: Session, project_id: str, entry_id: str, user: User) -> ProjectLegacyConsumption:
    entry = db.get(ProjectLegacyConsumption, entry_id)
    if not entry or entry.project_id != project_id:
        raise HTTPException(status_code=404, detail=translate("Lançamento de consumo anterior não encontrado", user.language))
    return entry


@router.get("/projects/{project_id}/legacy-consumption", response_model=list[ProjectLegacyConsumptionRead])
def list_legacy_consumption(
    project_id: str,
    user: User = Depends(require_roles(*MANAGEMENT_ROLES)),
    db: Session = Depends(get_db),
) -> list[dict]:
    _project_or_404(db, project_id, user)
    rows = db.scalars(
        select(ProjectLegacyConsumption)
        .where(ProjectLegacyConsumption.project_id == project_id)
        .order_by(ProjectLegacyConsumption.reference_date, ProjectLegacyConsumption.created_at)
    ).all()
    return [_read(r) for r in rows]


@router.post("/projects/{project_id}/legacy-consumption", response_model=ProjectLegacyConsumptionRead, status_code=201)
def create_legacy_consumption(
    project_id: str,
    data: ProjectLegacyConsumptionCreate,
    user: User = Depends(require_roles(*MANAGEMENT_ROLES)),
    db: Session = Depends(get_db),
) -> dict:
    _project_or_404(db, project_id, user)
    entry = ProjectLegacyConsumption(project_id=project_id, **data.model_dump())
    db.add(entry)
    db.flush()
    record_audit(
        db, entity_type="project", entity_id=project_id, action=AuditAction.UPDATE, user_id=user.id,
        details={"legacy_consumption_added": entry.id},
    )
    db.commit()
    db.refresh(entry)
    return _read(entry)


@router.patch("/projects/{project_id}/legacy-consumption/{entry_id}", response_model=ProjectLegacyConsumptionRead)
def update_legacy_consumption(
    project_id: str,
    entry_id: str,
    data: ProjectLegacyConsumptionUpdate,
    user: User = Depends(require_roles(*MANAGEMENT_ROLES)),
    db: Session = Depends(get_db),
) -> dict:
    _project_or_404(db, project_id, user)
    entry = _entry_or_404(db, project_id, entry_id, user)
    changes = data.model_dump(exclude_unset=True)
    for required in ("reference_date", "hours", "cost_per_hour"):
        if required in changes and changes[required] is None:
            raise HTTPException(status_code=422, detail=translate("Campo obrigatório não pode ser vazio", user.language))
    for field, value in changes.items():
        setattr(entry, field, value)
    record_audit(
        db, entity_type="project", entity_id=project_id, action=AuditAction.UPDATE, user_id=user.id,
        details={"legacy_consumption_updated": entry.id, "fields": sorted(changes.keys())},
    )
    db.commit()
    db.refresh(entry)
    return _read(entry)


@router.delete("/projects/{project_id}/legacy-consumption/{entry_id}", status_code=204)
def delete_legacy_consumption(
    project_id: str,
    entry_id: str,
    user: User = Depends(require_roles(*MANAGEMENT_ROLES)),
    db: Session = Depends(get_db),
) -> None:
    _project_or_404(db, project_id, user)
    entry = _entry_or_404(db, project_id, entry_id, user)
    db.delete(entry)
    record_audit(
        db, entity_type="project", entity_id=project_id, action=AuditAction.UPDATE, user_id=user.id,
        details={"legacy_consumption_removed": entry_id},
    )
    db.commit()
    return None


# ---------------------------------------------------------------------------
# Importação por planilha
# ---------------------------------------------------------------------------

_TEMPLATE_HEADERS = ["Código do projeto", "Data", "Horas", "Custo médio/hora", "Descrição"]


def _norm(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode()
    return "".join(ch for ch in text.lower() if ch.isalnum())


def _column_of(header: str) -> str | None:
    """Mapeia o cabeçalho (PT/ES, com ou sem acento) para o campo — None se
    não reconhecido (coluna ignorada)."""
    n = _norm(header)
    if n.startswith("codigo") or n in {"projeto", "proyecto", "projectcode", "project"}:
        return "project"
    if n in {"data", "fecha", "datareferencia", "fechareferencia", "periodo", "date"}:
        return "date"
    if n in {"horas", "hours"}:
        return "hours"
    if "custo" in n or "costo" in n or n.startswith("cost"):
        return "cost"
    if n.startswith("descri") or n.startswith("observ") or n == "notes":
        return "description"
    return None


def _parse_decimal(value: object) -> Decimal | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    text = str(value).strip().replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _parse_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool) and value > 59:
        try:
            return from_excel(value).date()
        except (ValueError, OverflowError):
            return None
    if isinstance(value, str):
        text = value.strip()
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
    return None


def build_import_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Consumo anterior"
    ws.append(_TEMPLATE_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.append(["0000000094", date(2026, 1, 31), 120.5, 42.5, "Horas consumidas até jan/2026 (sistema anterior)"])
    ws["B2"].number_format = "DD/MM/YYYY"
    for column, width in zip("ABCDE", (22, 14, 10, 20, 60)):
        ws.column_dimensions[column].width = width
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@router.get("/legacy-consumption/template.xlsx")
def legacy_consumption_template(user: User = Depends(require_roles(*MANAGEMENT_ROLES))) -> Response:
    return Response(
        content=build_import_template(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="modelo_consumo_anterior.xlsx"'},
    )


def _msg(lang: str, template: str, **kwargs: object) -> str:
    return translate(template, lang).format(**kwargs)


@router.post("/legacy-consumption/import", response_model=LegacyConsumptionImportResult, status_code=201)
async def import_legacy_consumption(
    file: UploadFile = File(...),
    user: User = Depends(require_roles(*MANAGEMENT_ROLES)),
    db: Session = Depends(get_db),
) -> dict:
    """Importa vários lançamentos de uma vez a partir de uma planilha .xlsx
    (modelo em GET /legacy-consumption/template.xlsx): colunas Código do
    projeto, Data, Horas, Custo médio/hora e Descrição (opcional). Tudo ou
    nada: se qualquer linha tiver erro, nada é gravado e a resposta 422 lista
    os problemas (com o número da linha na planilha)."""
    lang = user.language
    content = await file.read(MAX_IMPORT_BYTES + 1)
    if len(content) > MAX_IMPORT_BYTES:
        raise HTTPException(status_code=422, detail=translate("Arquivo muito grande (máximo 2 MB)", lang))
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        raw_rows = list(sheet.iter_rows(values_only=True))
    except Exception as exc:  # noqa: BLE001 — qualquer falha de leitura = arquivo inválido
        raise HTTPException(status_code=422, detail=translate("Arquivo inválido: envie uma planilha .xlsx", lang)) from exc
    if not raw_rows:
        raise HTTPException(status_code=422, detail=translate("A planilha está vazia", lang))

    columns: dict[str, int] = {}
    for index, header in enumerate(raw_rows[0]):
        field = _column_of(header) if header is not None else None
        if field and field not in columns:
            columns[field] = index
    missing = [name for name in ("project", "date", "hours", "cost") if name not in columns]
    if missing:
        raise HTTPException(
            status_code=422,
            detail=translate("Cabeçalho inválido: a planilha precisa das colunas Código do projeto, Data, Horas e Custo médio/hora", lang),
        )

    data_rows = [(number, row) for number, row in enumerate(raw_rows[1:], start=2) if any(cell not in (None, "") for cell in row)]
    if not data_rows:
        raise HTTPException(status_code=422, detail=translate("A planilha não tem nenhuma linha de dados", lang))
    if len(data_rows) > MAX_IMPORT_ROWS:
        raise HTTPException(status_code=422, detail=_msg(lang, "Máximo de {max} linhas por importação", max=MAX_IMPORT_ROWS))

    projects_by_code = {code.strip().upper(): pid for pid, code in db.execute(select(Project.id, Project.code)).all()}
    errors: list[str] = []
    parsed: list[tuple[str, date, Decimal, Decimal, str | None]] = []

    def cell(row: tuple, field: str):
        index = columns.get(field)
        return row[index] if index is not None and index < len(row) else None

    for number, row in data_rows:
        code = str(cell(row, "project") or "").strip()
        project_id = projects_by_code.get(code.upper()) if code else None
        reference_date = _parse_date(cell(row, "date"))
        hours = _parse_decimal(cell(row, "hours"))
        cost = _parse_decimal(cell(row, "cost"))
        description = str(cell(row, "description") or "").strip() or None
        problems: list[str] = []
        if not code:
            problems.append(translate("código do projeto vazio", lang))
        elif project_id is None:
            problems.append(_msg(lang, 'projeto "{code}" não encontrado', code=code))
        if reference_date is None:
            problems.append(translate("data inválida", lang))
        if hours is None or hours <= 0:
            problems.append(translate("horas devem ser um número maior que zero", lang))
        if cost is None or cost < 0:
            problems.append(translate("custo médio/hora deve ser um número maior ou igual a zero", lang))
        if description and len(description) > 255:
            problems.append(translate("descrição acima de 255 caracteres", lang))
        if problems:
            errors.append(_msg(lang, "linha {row}: {problems}", row=number, problems=", ".join(problems)))
        else:
            parsed.append((project_id, reference_date, hours.quantize(_CENT), cost.quantize(_CENT), description))  # type: ignore[arg-type]

    if errors:
        shown = "; ".join(errors[:MAX_REPORTED_ERRORS])
        extra = len(errors) - MAX_REPORTED_ERRORS
        suffix = _msg(lang, " (+{n} outros erros)", n=extra) if extra > 0 else ""
        raise HTTPException(
            status_code=422,
            detail=_msg(lang, "Importação cancelada, nada foi gravado — {n} linha(s) com erro: ", n=len(errors)) + shown + suffix,
        )

    per_project: dict[str, int] = defaultdict(int)
    total_hours = Decimal("0")
    total_cost = Decimal("0")
    for project_id, reference_date, hours, cost, description in parsed:
        db.add(
            ProjectLegacyConsumption(
                project_id=project_id, reference_date=reference_date, hours=hours, cost_per_hour=cost, description=description
            )
        )
        per_project[project_id] += 1
        total_hours += hours
        total_cost += hours * cost
    for project_id, count in per_project.items():
        record_audit(
            db, entity_type="project", entity_id=project_id, action=AuditAction.UPDATE, user_id=user.id,
            details={"legacy_consumption_imported": count},
        )
    db.commit()
    return {
        "created": len(parsed),
        "projects": len(per_project),
        "total_hours": total_hours.quantize(_CENT),
        "total_cost": total_cost.quantize(_CENT),
    }
